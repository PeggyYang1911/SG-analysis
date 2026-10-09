"""
Styled SG+ % line figure for the RPMI 8226 panel - 2 cell lines.

Same look and pipeline as rpmi_4celllines_styled.py (Okabe-Ito colours, legend
above the axes, mean +/- SD with jittered per-image points), but restricted to
the Parental and HDAC6 Knock-out lines only. The BTZ-resistant and
BTZ-resistant HDAC6 KO lines are dropped.

Counting rule: unique Cell_ID (0 = background, excluded); per-cell punctae =
SUM of Punctae_Count over that cell's slices; cell is SG+ if that sum > 0.

------------------------------------------------------------------------------
HOW TO CHANGE THE FIGURE
------------------------------------------------------------------------------
1) Edit the CONFIG block below (axis limits, ticks, legend position, fonts,
   error bars SD/SEM, colours, labels) and re-run:
       python rpmi_2celllines_styled.py

2) Drag the legend with the mouse instead of guessing coordinates:
       python rpmi_2celllines_styled.py --show
   An interactive window opens. Drag the legend wherever you like, use the
   toolbar to pan/zoom the axes, then use the toolbar save (disk) icon.
   When you close the window the script prints the legend position and the
   axis limits it ended up at, ready to paste into the CONFIG block.

3) Edit in Illustrator / Inkscape / PowerPoint: every run also writes .svg and
   .pdf next to the .png, with real (non-outlined) text, so axis numbers,
   labels and the legend stay editable as text objects.
------------------------------------------------------------------------------

Outputs into OUT_DIR:
  Figure_SGpositive_lines_2lines.png / .svg / .pdf
  SG_lines_summary_2lines.xlsx  (Per_image + Group_means)
"""
import csv, glob, os, re, sys
import numpy as np
import matplotlib
if "--show" not in sys.argv:
    matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl

from sg_provenance import (PER_CELL_TAIL, cell_rows, parse_dates, slices_used,
                           write_sheet)

# ============================== DATA SOURCES ================================
BASE_5PER = (r"J:\FATHO_LAB\YANG Pei-Tzu\Biological evaluation\Immunostaining"
             r"\260624_PY_260611_RPMI_KO_R_AS\v15 analysis\5slices_subset"
             r"\raw_data_5perGroup")

OUT_DIR = (r"J:\FATHO_LAB\YANG Pei-Tzu\Biological evaluation\Immunostaining"
           r"\260624_PY_260611_RPMI_KO_R_AS\v15 analysis\5slices_subset")

# (legend label, colour, folder, include-substring, exclude-substrings)
LINES = [
    ("Parental",               "#0072B2", BASE_5PER, "_RPMI_AS_",      ()),
    ("HDAC6 Knock-out",        "#E69F00", BASE_5PER, "_RPMI_KO_AS_",   ()),
]
TIMEPOINTS = [0, 15, 30, 45]

# ================================= CONFIG ===================================
# ---- axes ------------------------------------------------------------------
XLABEL      = "1 mM NaAsO$_2$ (min)"    # $_2$ = subscript
YLABEL      = "SG+ cells (%)"
XLIM        = (-3, 48)                  # None = auto
YLIM        = (-2, 101)                 # None = auto
XTICKS      = [0, 15, 30, 45]           # None = auto
YTICKS      = [0, 20, 40, 60, 80, 100]  # None = auto

# ---- legend ----------------------------------------------------------------
LEGEND_TITLE  = "RPMI 8226"
LEGEND_SHOW   = True
LEGEND_LOC    = "lower left"    # which corner OF THE LEGEND BOX is anchored
LEGEND_BBOX   = (0.0, 1.01)     # where that corner sits, in axes coordinates:
                                # (0,0) = bottom-left of the plot area,
                                # (1,1) = top-right; values >1 or <0 place the
                                # legend outside the axes. Use None to let
                                # matplotlib place it from LEGEND_LOC alone
                                # (e.g. LEGEND_LOC="upper left" -> inside).
LEGEND_NCOL   = 1
LEGEND_FRAME  = False           # True = draw a box around the legend

# ---- sizes / fonts ---------------------------------------------------------
FIGSIZE       = (8.4, 9.0)      # inches (was 6.6 high; taller = longer y-axis,
                                # font sizes below are in points so text is
                                # unaffected - scale them too if you want the
                                # text to keep the same relative size)
FONT_FAMILY   = "DejaVu Sans"
FONT_BASE     = 22
FONT_AXLABEL  = 26
FONT_TICKS    = 24
FONT_LEGEND   = 21
FONT_LEGTITLE = 22

# ---- marks -----------------------------------------------------------------
ERRORBAR      = "SD"            # "SD" or "SEM"
LINEWIDTH     = 3.0
MARKERSIZE    = 11
CAPSIZE       = 4
SHOW_POINTS   = True            # jittered individual images
POINT_SIZE    = 40
POINT_ALPHA   = 0.35
POINT_JITTER  = 1.2             # in x-units (minutes)
SPINE_WIDTH   = 1.6
HIDE_TOP_RIGHT_SPINES = True

# ---- output ----------------------------------------------------------------
OUT_STEM  = "Figure_SGpositive_lines_2lines"
FORMATS   = ("png", "svg", "pdf")
DPI       = 200
OUT_XLSX  = "SG_lines_summary_2lines.xlsx"

# Extra SVG for PowerPoint: same figure but drawn in a font Windows actually
# has (DejaVu Sans ships with matplotlib and is NOT installed on this PC, so
# PowerPoint would substitute it and shift the text). Insert this file in
# PowerPoint and use Graphics Format > Convert to Shape to get editable text.
PPT_SVG   = True
PPT_FONT  = "Arial"
PPT_SUFFIX = "_PPT"
# ============================================================================


def is_data(row, ci):
    try:
        int(float(row[ci])); return True
    except (ValueError, IndexError):
        return False


def image_stats(path):
    """-> (image, SG+%, SG/cell, SG/SG+cell, cells, slices, n_slices) or None."""
    with open(path, newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return None
    h = rows[0]
    try:
        ci = h.index("Cell_ID"); pi = h.index("Punctae_Count")
        si = h.index("Slice")
    except ValueError:
        return None
    data = [r for r in rows[1:] if is_data(r, ci)]
    if not data:
        return None
    cellp = {}
    cell_sl = {}
    for r in data:
        cid = r[ci].strip()
        if cid == "0":
            continue
        cellp[cid] = cellp.get(cid, 0) + int(float(r[pi]))
        cell_sl.setdefault(cid, set()).add(r[si].strip())
    if not cellp:
        return None
    ncell = len(cellp)
    withp = sum(1 for v in cellp.values() if v > 0)
    tot = sum(cellp.values())
    sl_str, n_sl = slices_used(cell_sl)
    return (data[0][0], withp / ncell * 100, tot / ncell,
            tot / withp if withp else 0.0,
            cell_rows(cellp, cell_sl), sl_str, n_sl)


def collect():
    # (label, tp, image, SG+%, SG/cell, SG/SG+cell, cells, slices, n_slices,
    #  source_csv, acq_date, ana_date) -- vals() indexes 3/4/5, so the metrics
    #  must keep their positions; provenance is appended after them.
    recs = []
    for label, _c, folder, sub, excl in LINES:
        for path in sorted(glob.glob(os.path.join(folder, "*.csv"))):
            fn = os.path.basename(path)
            if sub not in fn or any(e in fn for e in excl):
                continue
            m = re.search(r"(\d+)\s*_?\s*min", fn)
            if not m or int(m.group(1)) not in TIMEPOINTS:
                continue
            st = image_stats(path)
            if st is None:
                print("  skipped (no data):", fn)
                continue
            ana_date, acq_date = parse_dates(fn)
            recs.append((label, int(m.group(1))) + st + (fn, acq_date, ana_date))
    return recs


def vals(recs, label, tp, idx):
    return np.array([r[idx] for r in recs if r[0] == label and r[1] == tp])


def build_figure(recs, font_family=None):
    fam = font_family or FONT_FAMILY
    # keep text as text in svg/pdf so it stays editable in Illustrator/Inkscape
    plt.rcParams.update({"font.size": FONT_BASE,
                         "font.family": fam,
                         # render the NaAsO2 subscript in the same font, so the
                         # file references only one font family
                         "mathtext.fontset": "custom", "mathtext.rm": fam,
                         "mathtext.it": f"{fam}:italic", "mathtext.bf": f"{fam}:bold",
                         "svg.fonttype": "none", "pdf.fonttype": 42,
                         "ps.fonttype": 42,
                         # so the toolbar save button in --show mode produces
                         # the same quality/margins as the script output
                         "savefig.dpi": DPI, "savefig.bbox": "tight"})
    fig, ax = plt.subplots(figsize=FIGSIZE)
    rng = np.random.default_rng(0)
    for label, color, *_ in LINES:
        xs, means, errs = [], [], []
        for tp in TIMEPOINTS:
            v = vals(recs, label, tp, 3)
            if not len(v):
                continue
            sd = v.std(ddof=1) if len(v) > 1 else 0.0
            xs.append(tp)
            means.append(v.mean())
            errs.append(sd / np.sqrt(len(v)) if ERRORBAR.upper() == "SEM" else sd)
            if SHOW_POINTS:
                jit = rng.uniform(-POINT_JITTER, POINT_JITTER, size=len(v))
                ax.scatter(np.array([tp] * len(v)) + jit, v, color=color,
                           s=POINT_SIZE, alpha=POINT_ALPHA, edgecolors="none",
                           zorder=2)
        ax.errorbar(xs, means, yerr=errs, marker="o", ms=MARKERSIZE,
                    lw=LINEWIDTH, color=color, capsize=CAPSIZE, capthick=2.2,
                    label=label, zorder=3)

    ax.set_xlabel(XLABEL, fontsize=FONT_AXLABEL)
    ax.set_ylabel(YLABEL, fontsize=FONT_AXLABEL)
    if XTICKS is not None:
        ax.set_xticks(XTICKS)
    if YTICKS is not None:
        ax.set_yticks(YTICKS)
    if XLIM is not None:
        ax.set_xlim(*XLIM)
    if YLIM is not None:
        ax.set_ylim(*YLIM)
    ax.tick_params(labelsize=FONT_TICKS, width=SPINE_WIDTH, length=7)
    if HIDE_TOP_RIGHT_SPINES:
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_linewidth(SPINE_WIDTH)

    leg = None
    if LEGEND_SHOW:
        leg = ax.legend(title=LEGEND_TITLE, loc=LEGEND_LOC,
                        bbox_to_anchor=LEGEND_BBOX, ncol=LEGEND_NCOL,
                        fontsize=FONT_LEGEND, title_fontsize=FONT_LEGTITLE,
                        frameon=LEGEND_FRAME, handlelength=1.4,
                        borderaxespad=0.0)
        leg._legend_box.align = "left"
        leg.set_draggable(True)   # drag it around in the --show window
    fig.tight_layout()
    return fig, ax, leg


def save_figure(fig):
    outs = []
    for ext in FORMATS:
        out = os.path.join(OUT_DIR, f"{OUT_STEM}.{ext}")
        fig.savefig(out, dpi=DPI, bbox_inches="tight")
        outs.append(out)
    return outs


def workbook(recs):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Per_image"
    ws.append(["Cell_line", "Timepoint_min", "Image", "Source_CSV",
               "Acquisition_date", "Analysis_date", "Slices_used", "N_slices",
               "SG_positive_(%)", "SG_per_cell", "SG_per_SG+cell"])
    order = [l[0] for l in LINES]
    ordered = sorted(recs, key=lambda r: (order.index(r[0]), r[1], r[2]))
    for r in ordered:
        ws.append([r[0], r[1], r[2], r[9], r[10], r[11], r[7], r[8],
                   round(r[3], 3), round(r[4], 4), round(r[5], 4)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([24, 14, 46, 46, 17, 15, 16, 10, 14, 14, 16], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    write_sheet(
        wb.create_sheet("Per_cell_counts"),
        [("Cell_line", "line", 24), ("Timepoint_min", "tp", 14),
         ("Image", "image", 46), ("Source_CSV", "csv", 46)] + PER_CELL_TAIL,
        [dict(c, line=r[0], tp=r[1], image=r[2], csv=r[9],
              acq_date=r[10], ana_date=r[11])
         for r in ordered for c in r[6]])
    ws2 = wb.create_sheet("Group_means")
    ws2.append(["Cell_line", "Timepoint_min", "n", "Mean_SG+_(%)", "SD_SG+_(%)",
                "SEM_SG+_(%)", "Mean_SG_per_cell", "SD_SG_per_cell",
                "Mean_SG_per_SG+cell", "SD_SG_per_SG+cell"])
    for label in order:
        for tp in TIMEPOINTS:
            v = vals(recs, label, tp, 3)
            if not len(v):
                continue
            def ms(idx):
                a = vals(recs, label, tp, idx)
                return (round(a.mean(), 4),
                        round(a.std(ddof=1), 4) if len(a) > 1 else 0.0)
            m3, s3 = ms(3); m4, s4 = ms(4); m5, s5 = ms(5)
            ws2.append([label, tp, len(v), m3, s3,
                        round(s3 / np.sqrt(len(v)), 4) if len(v) > 1 else 0.0,
                        m4, s4, m5, s5])
    out = os.path.join(OUT_DIR, OUT_XLSX)
    wb.save(out)
    return out


def main():
    show = "--show" in sys.argv
    recs = collect()
    print(f"{'cell line':26}" + "".join(f"{str(t)+'min':>12}" for t in TIMEPOINTS))
    for label, *_ in LINES:
        cells = []
        for tp in TIMEPOINTS:
            v = vals(recs, label, tp, 3)
            cells.append(f"{v.mean():6.1f}(n={len(v)})" if len(v) else "     -")
        print(f"{label:26}" + "".join(f"{c:>12}" for c in cells))

    fig, ax, leg = build_figure(recs)

    if show:
        print("\nInteractive mode:")
        print("  - drag the legend with the mouse")
        print("  - toolbar: pan/zoom the axes, disk icon = save what you see")
        print("  - close the window to print the final settings\n")
        plt.show()
        if leg is not None:
            bb = leg.get_window_extent().transformed(ax.transAxes.inverted())
            print("Paste into the CONFIG block:")
            print('  LEGEND_LOC  = "lower left"')
            print(f"  LEGEND_BBOX = ({bb.x0:.3f}, {bb.y0:.3f})")
        print(f"  XLIM        = ({ax.get_xlim()[0]:.2f}, {ax.get_xlim()[1]:.2f})")
        print(f"  YLIM        = ({ax.get_ylim()[0]:.2f}, {ax.get_ylim()[1]:.2f})")
        return

    for out in save_figure(fig):
        print("Saved:", out)
    plt.close(fig)

    if PPT_SVG:
        ppt_fig, _ax, _leg = build_figure(recs, font_family=PPT_FONT)
        out = os.path.join(OUT_DIR, f"{OUT_STEM}{PPT_SUFFIX}.svg")
        ppt_fig.savefig(out, dpi=DPI, bbox_inches="tight")
        plt.close(ppt_fig)
        print(f"Saved: {out}   ({PPT_FONT}, for PowerPoint)")

    print("Saved:", workbook(recs))


if __name__ == "__main__":
    main()
