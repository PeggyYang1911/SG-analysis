"""
Shared provenance helpers for the stress-granule analysis scripts.

Every script in this folder that writes counts must record WHERE each number came
from: the source CSV, the imaging and analysis dates, the Series, the Cell_ID and
the z-slices behind it. That keeps any figure traceable back to the images it was
built from, and lets per-cell rows from different experiments be concatenated into
one pooled table later.

`Source_CSV` is the join key. It is unique per image and is present in every
workbook these scripts produce, so it is the column to merge on when pooling.

Imported by sg_analysis.py, celline_compare.py, celline_compare_5slices.py,
combined_lines.py and rpmi_4celllines_styled.py. They all live in this folder and
are run by full path, so `import sg_provenance` resolves through sys.path[0].

This module deliberately does NOT do any cleaning or counting -- each script keeps
its own cleaning rules (which differ: in-memory clean, 5-slice subset, or none at
all for pre-subset folders), so adding provenance can never change a number.
"""
import re

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATE_RE = re.compile(r"(?<!\d)(\d{6})(?!\d)")   # YYMMDD tokens

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="4F81BD")

# Trailing columns shared by every per-cell sheet: (header, record key, width).
# Each script prepends its own grouping columns (cell line, timepoint, image, ...).
PER_CELL_TAIL = [
    ("Acquisition_date", "acq_date", 17),
    ("Analysis_date", "ana_date", 15),
    ("Cell_ID", "cell_id", 9),
    ("Slices", "slices", 16),
    ("N_slices", "n_slices", 10),
    ("Punctae_per_cell", "punctae", 18),
    ("SG_positive", "sg_pos", 12),
]


def parse_dates(fname):
    """Return (analysis_date, acquisition_date) as YYYY-MM-DD strings.

    File names carry leading YYMMDD tokens: the analysis date first, then the
    imaging date when the two differ ("260624_260617_PY_RPMI_AS_0min_...").  With
    a single token both dates are the same; with none, both come back empty.
    Tokens that are not a valid month/day are ignored, and the pattern requires
    exactly six digits, so "Series001" and concentrations like "1000nM" can never
    be mistaken for a date.
    """
    found = []
    for tok in DATE_RE.findall(fname):
        mm, dd = int(tok[2:4]), int(tok[4:6])
        if 1 <= mm <= 12 and 1 <= dd <= 31:
            found.append("20%s-%s-%s" % (tok[:2], tok[2:4], tok[4:6]))
        if len(found) == 2:
            break
    if not found:
        return "", ""
    return found[0], (found[1] if len(found) > 1 else found[0])


def slice_sort(s):
    """Sort z-slices numerically, pushing anything unparseable to the end."""
    try:
        return int(float(s))
    except ValueError:
        return 10 ** 9


def cell_rows(cell_punct, cell_slices):
    """One dict per counted cell, ordered by Cell_ID.

    cell_punct : {Cell_ID -> summed Punctae_Count}
    cell_slices: {Cell_ID -> set of slices that cell was measured in}
    """
    out = []
    for cid in sorted(cell_punct, key=lambda c: int(float(c))):
        sl = sorted(cell_slices[cid], key=slice_sort)
        v = cell_punct[cid]
        out.append({
            "cell_id": int(float(cid)),
            "slices": ",".join(sl),
            "n_slices": len(sl),
            "punctae": v,
            "sg_pos": "Yes" if v > 0 else "No",
        })
    return out


def slices_used(cell_slices):
    """(comma-separated slice list, count) across every cell in one image."""
    sl = sorted({s for ss in cell_slices.values() for s in ss}, key=slice_sort)
    return ",".join(sl), len(sl)


def write_sheet(ws, cols, records, autofilter=True):
    """Fill a worksheet from `cols` = [(header, record key, width)] and dicts."""
    ws.append([c[0] for c in cols])
    for d in records:
        ws.append([d.get(c[1]) for c in cols])
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = c[2]
    ws.freeze_panes = "A2"
    if autofilter:
        ws.auto_filter.ref = ws.dimensions
    return ws
