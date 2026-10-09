"""
Line figures + statistics from a hand-selected subset sheet in
ThreeCellLine_comparison.xlsx (e.g. "3 sets per group": 3 chosen Series per
cell line x timepoint).

Reads the metrics already present in the sheet (SG_positive_(%),
Mean_punctae_per_cell, Mean_punctae_per_SG+_cell) -- it does NOT recompute from
CSVs, so it exactly matches the rows the user selected.

Outputs (into the same folder as the workbook), WITHOUT touching the source
workbook (openpyxl resave would drop its embedded figures):
  Fig_line_3sets_SGpositive_percent.png
  Fig_line_3sets_SG_per_cell.png
  Fig_line_3sets_SG_per_SGpositive.png
  ThreeCellLine_3sets_selected.xlsx  (Selected_data, Group_means, Statistics, Figures)

Stat unit = per image/Series. At each timepoint: one-way ANOVA across the cell
lines for each metric (n>=2 groups). Line plots show mean +/- SEM.
"""
import os
import numpy as np
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FOLDER = r"J:\FATHO_LAB\YANG Pei-Tzu\Biological evaluation\Immunostaining\260601_PY_260522 Sterallis KMS WT and KO AS\v15 analysis\3celline_comparison_WT_KO_R"
WORKBOOK = os.path.join(FOLDER, "ThreeCellLine_comparison.xlsx")
SHEET = "3 sets per group"

CELL_LINES = [("KMS WT", "#4C72B0"), ("KMS KO", "#C44E52"), ("KMS-R", "#55A868")]
TIMEPOINTS = [0, 15, 30, 45]

# metric key -> (column header in the sheet, y-label, title stub, file stub)
METRICS = [
    ("SG_positive_(%)", "SG+ cells (%)", "Stress-granule-positive cells (%)",
     "Fig_line_3sets_SGpositive_percent.png"),
    ("Mean_punctae_per_cell", "Number of SG per cell", "SG per cell (all cells)",
     "Fig_line_3sets_SG_per_cell.png"),
    ("Mean_punctae_per_SG+_cell", "SG per SG+ cell", "SG per SG+ cell",
     "Fig_line_3sets_SG_per_SGpositive.png"),
]


def load_rows():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    data = []
    for r in rows[1:]:
        if r is None or r[idx["Cell_line"]] is None:
            continue
        rec = {
            "line": str(r[idx["Cell_line"]]).strip(),
            "tp": int(r[idx["Timepoint_min"]]),
            "image": r[idx["Image"]],
            "series": r[idx["Series"]],
        }
        for key, _, _, _ in METRICS:
            rec[key] = float(r[idx[key]])
        data.append(rec)
    return data, idx


def vals(data, line, tp, key):
    return [d[key] for d in data if d["line"] == line and d["tp"] == tp]


def line_plot(data, key, ylabel, title, outname):
    fig, ax = plt.subplots(figsize=(7.5, 5.5))
    rng = np.random.default_rng(0)
    for line, color in CELL_LINES:
        xs, means, sems = [], [], []
        for tp in TIMEPOINTS:
            v = vals(data, line, tp, key)
            if v:
                xs.append(tp)
                means.append(np.mean(v))
                sems.append(np.std(v, ddof=1) / np.sqrt(len(v)) if len(v) > 1 else 0.0)
                jit = rng.uniform(-0.8, 0.8, size=len(v))
                ax.scatter(np.array([tp] * len(v)) + jit, v, color=color,
                           s=14, alpha=0.5, zorder=2)
        ax.errorbar(xs, means, yerr=sems, marker="o", ms=7, lw=2.2,
                    color=color, capsize=3, label=line, zorder=3)
    ax.set_xlabel("Sodium arsenite exposure (min)")
    ax.set_ylabel(ylabel)
    ax.set_title(title + "  (n=3/group; mean +/- SEM)")
    ax.set_xticks(TIMEPOINTS)
    ax.legend(frameon=False)
    fig.tight_layout()
    out = os.path.join(FOLDER, outname)
    fig.savefig(out, dpi=200)
    plt.close(fig)
    return out


def sig_mark(p):
    if not np.isfinite(p):
        return "n/a"
    if p < 0.001:
        return "***"
    if p < 0.01:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


def main():
    data, _ = load_rows()
    print(f"Loaded {len(data)} rows from '{SHEET}'")
    for line, _ in CELL_LINES:
        cnt = {tp: len(vals(data, line, tp, "SG_positive_(%)")) for tp in TIMEPOINTS}
        print(f"  {line}: " + ", ".join(f"{tp}min n={cnt[tp]}" for tp in TIMEPOINTS))

    figs = {}
    for key, ylabel, title, outname in METRICS:
        figs[key] = line_plot(data, key, ylabel, title, outname)

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="404040")

    def style(ws, ncol, row=1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

    # Selected_data
    ws = wb.active
    ws.title = "Selected_data"
    cols = ["Cell_line", "Timepoint_min", "Image", "Series",
            "SG_positive_(%)", "Mean_punctae_per_cell", "Mean_punctae_per_SG+_cell"]
    ws.append(cols)
    order = {lab: i for i, (lab, _) in enumerate(CELL_LINES)}
    for d in sorted(data, key=lambda d: (order.get(d["line"], 9), d["tp"], d["series"])):
        ws.append([d["line"], d["tp"], d["image"], d["series"],
                   round(d["SG_positive_(%)"], 3),
                   round(d["Mean_punctae_per_cell"], 4),
                   round(d["Mean_punctae_per_SG+_cell"], 4)])
    style(ws, len(cols))
    ws.freeze_panes = "A2"
    for i, w in enumerate([12, 14, 46, 8, 16, 22, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Group_means
    ws2 = wb.create_sheet("Group_means")
    ws2.append(["Cell_line", "Timepoint_min", "n",
                "Mean_SG+_(%)", "SEM_SG+_(%)",
                "Mean_SG_per_cell", "SEM_SG_per_cell",
                "Mean_SG_per_SG+cell", "SEM_SG_per_SG+cell"])

    def msem(v):
        m = float(np.mean(v))
        s = float(np.std(v, ddof=1) / np.sqrt(len(v))) if len(v) > 1 else 0.0
        return round(m, 4), round(s, 4)

    for line, _ in CELL_LINES:
        for tp in TIMEPOINTS:
            v1 = vals(data, line, tp, "SG_positive_(%)")
            if not v1:
                continue
            m1, s1 = msem(v1)
            m2, s2 = msem(vals(data, line, tp, "Mean_punctae_per_cell"))
            m3, s3 = msem(vals(data, line, tp, "Mean_punctae_per_SG+_cell"))
            ws2.append([line, tp, len(v1), m1, s1, m2, s2, m3, s3])
    style(ws2, 9)
    for i, w in enumerate([12, 14, 6, 14, 14, 18, 18, 20, 20], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Statistics: per-timepoint one-way ANOVA across cell lines
    ws3 = wb.create_sheet("Statistics")
    ws3.cell(1, 1, "One-way ANOVA across cell lines at each timepoint "
                   "(selected 3 series/group; per-image)").font = Font(bold=True)
    ws3.append(["Metric", "Timepoint_min", "cell_lines_(n)", "F", "p_value", "Significance"])
    style(ws3, 6, row=2)
    printout = {}
    for key, ylabel, *_ in METRICS:
        printout[ylabel] = []
        for tp in TIMEPOINTS:
            groups, tags = [], []
            for line, _ in CELL_LINES:
                v = vals(data, line, tp, key)
                if len(v) >= 2:
                    groups.append(np.array(v))
                    tags.append(f"{line}({len(v)})")
            if len(groups) >= 2:
                F, p = stats.f_oneway(*groups)
            else:
                F, p = float("nan"), float("nan")
            ws3.append([ylabel, tp, ", ".join(tags),
                        round(float(F), 4) if np.isfinite(F) else "n/a",
                        round(float(p), 6) if np.isfinite(p) else "n/a",
                        sig_mark(p)])
            if np.isfinite(p):
                printout[ylabel].append(f"{tp}min p={p:.4g}({sig_mark(p)})")
        ws3.append([])
    ws3.append(["Significance: *** p<0.001, ** p<0.01, * p<0.05, ns = not significant"])
    for i, w in enumerate([24, 14, 34, 12, 12, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Figures
    ws4 = wb.create_sheet("Figures")
    for (key, *_), row in zip(METRICS, [1, 30, 59]):
        ws4.add_image(XLImage(figs[key]), f"A{row}")

    out_xlsx = os.path.join(FOLDER, "ThreeCellLine_3sets_selected.xlsx")
    wb.save(out_xlsx)

    print("\n=== Per-timepoint ANOVA across cell lines (selected 3/group) ===")
    for ylabel, line in printout.items():
        print(f"  {ylabel}: " + ", ".join(line))
    print("\nSaved:")
    for f in [figs[k] for k, *_ in METRICS] + [out_xlsx]:
        print("  ", f)


if __name__ == "__main__":
    main()
