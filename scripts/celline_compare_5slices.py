"""
Three-cell-line SG comparison (KMS WT / KMS KO / KMS-R) over the arsenite time
course -- 5-SLICE-SUBSET version.

Same full image set and sources as celline_compare.py (the Summary_per_image /
"1st sheet" scope), but for every image, after the standard in-memory cleaning
(drop Cell_ID==0, drop slices with <5 cells) it keeps only the N_SLICES z-slices
with the HIGHEST cell coverage (distinct Cell_IDs; ties -> lowest slice number),
exactly like slice_subset.py. Metrics are then computed on those slices only.
Nothing on disk is modified.

Outputs into OUT_DIR (grouped bars + line plots per metric + workbook).
"""
import csv
import glob
import os
import re
import numpy as np
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from sg_provenance import (PER_CELL_TAIL, cell_rows, parse_dates, slices_used,
                           write_sheet)

BASE = r"J:\FATHO_LAB\YANG Pei-Tzu\Biological evaluation\Immunostaining"
WT_DIR = os.path.join(BASE, r"260601_PY_260522 Sterallis KMS WT and KO AS\v15 analysis\KMS_WT")
KO_DIR = os.path.join(BASE, r"260601_PY_260522 Sterallis KMS WT and KO AS\v15 analysis\KMS_KO")
R_DIR  = os.path.join(BASE, r"260612_PY_260609 Sterallis KMS -R AS BTZ\v15 analysis\KMS-R AS timecourse analysis")

CELL_LINES = [
    ("KMS WT",  "#4C72B0", [WT_DIR]),
    ("KMS KO",  "#C44E52", [KO_DIR]),
    ("KMS-R",   "#55A868", [R_DIR]),
]
TIMEPOINTS = [0, 15, 30, 45]
MIN_CELLS_PER_SLICE = 5
N_SLICES = 5
OUT_DIR = os.path.join(BASE, r"260601_PY_260522 Sterallis KMS WT and KO AS\v15 analysis\3celline_comparison_WT_KO_R_5slices")

TIME_RE = re.compile(r"(\d+)\s*_?\s*min", re.IGNORECASE)
SERIES_RE = re.compile(r"Series(\d+)")


def timepoint(fname):
    m = TIME_RE.search(fname)
    return int(m.group(1)) if m else None


def image_metrics(path):
    with open(path, "r", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return None
    header = rows[0]
    try:
        s_i = header.index("Slice")
        c_i = header.index("Cell_ID")
        p_i = header.index("Punctae_Count")
    except ValueError:
        return None

    def is_data(r):
        if len(r) < len(header):
            return False
        try:
            int(float(r[c_i]))
            return True
        except (ValueError, IndexError):
            return False

    data = [r for r in rows[1:] if is_data(r)]
    if not data:
        return None
    # standard cleaning
    data = [r for r in data if r[c_i].strip() != "0"]
    counts = {}
    for r in data:
        counts[r[s_i]] = counts.get(r[s_i], 0) + 1
    valid = {s for s, c in counts.items() if c >= MIN_CELLS_PER_SLICE}
    data = [r for r in data if r[s_i] in valid]
    if not data:
        return None
    # 5-slice subset: keep top-N slices by distinct Cell_ID coverage
    perslice = {}
    for r in data:
        perslice.setdefault(r[s_i], set()).add(r[c_i])
    ranked = sorted(perslice.items(), key=lambda kv: (-len(kv[1]), int(kv[0])))
    chosen = {s for s, _ in ranked[:N_SLICES]}
    data = [r for r in data if r[s_i] in chosen]
    n_slices_used = len(chosen)
    # aggregate per unique Cell_ID, remembering which slices each cell came from
    cellp = {}
    cell_sl = {}
    for r in data:
        cid = r[c_i].strip()
        cellp[cid] = cellp.get(cid, 0) + int(float(r[p_i]))
        cell_sl.setdefault(cid, set()).add(r[s_i].strip())
    percell = list(cellp.values())
    if not percell:
        return None
    total = len(percell)
    withp = sum(1 for v in percell if v > 0)
    tot_p = sum(percell)
    return {
        "total_cells": total,
        "cells_with": withp,
        "pct": withp / total * 100.0,
        "mean_per_cell": tot_p / total,
        "mean_per_pos": (tot_p / withp) if withp else 0.0,
        "total_punctae": tot_p,
        "n_slices": n_slices_used,
        "source_image": data[0][0],
        "slices_str": slices_used(cell_sl)[0],
        "cells": cell_rows(cellp, cell_sl),
    }


def collect():
    per_image = []
    for label, color, dirs in CELL_LINES:
        for d in dirs:
            for path in sorted(glob.glob(os.path.join(d, "*.csv"))):
                fname = os.path.basename(path)
                tp = timepoint(fname)
                if tp is None or tp not in TIMEPOINTS:
                    continue
                m = image_metrics(path)
                if m is None:
                    print(f"  skip (no data): {fname}")
                    continue
                sm = SERIES_RE.search(fname)
                ana_date, acq_date = parse_dates(fname)
                m.update(line=label, tp=tp, image=fname, csv=fname,
                         acq_date=acq_date, ana_date=ana_date,
                         series=int(sm.group(1)) if sm else -1)
                per_image.append(m)
    return per_image


def vals(per_image, label, tp, key):
    return [d[key] for d in per_image if d["line"] == label and d["tp"] == tp]


def grouped_bar(per_image, key, ylabel, title, outname):
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(8.5, 5.5))
    n = len(CELL_LINES)
    width = 0.8 / n
    x = np.arange(len(TIMEPOINTS))
    for li, (label, color, _) in enumerate(CELL_LINES):
        means, sds = [], []
        for tp in TIMEPOINTS:
            v = vals(per_image, label, tp, key)
            means.append(np.mean(v) if v else 0.0)
            sds.append(np.std(v, ddof=1) if len(v) > 1 else 0.0)
        off = (li - (n - 1) / 2) * width
        ax.bar(x + off, means, width, yerr=sds, capsize=3, color=color,
               edgecolor="black", alpha=0.85, label=label, zorder=2)
        for xi, tp in enumerate(TIMEPOINTS):
            v = vals(per_image, label, tp, key)
            jit = rng.uniform(-width * 0.3, width * 0.3, size=len(v))
            ax.scatter(x[xi] + off + jit, v, color="black", s=10, zorder=3)
    ax.set_xticks(x)
    ax.set_xticklabels([f"{t} min" for t in TIMEPOINTS])
    ax.set_xlabel("Sodium arsenite exposure")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.legend(frameon=False)
    fig.tight_layout()
    out = os.path.join(OUT_DIR, outname)
    fig.savefig(out, dpi=200)
    plt.close(fig)
    return out


def line_plot(per_image, key, ylabel, title, outname):
    fig, ax = plt.subplots(figsize=(7.5, 5.5))
    for label, color, _ in CELL_LINES:
        xs, means, sems = [], [], []
        for tp in TIMEPOINTS:
            v = vals(per_image, label, tp, key)
            if v:
                xs.append(tp)
                means.append(np.mean(v))
                sems.append(np.std(v, ddof=1) / np.sqrt(len(v)) if len(v) > 1 else 0.0)
        ax.errorbar(xs, means, yerr=sems, marker="o", ms=6, lw=2,
                    color=color, capsize=3, label=label)
    ax.set_xlabel("Sodium arsenite exposure (min)")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.set_xticks(TIMEPOINTS)
    ax.legend(frameon=False)
    fig.tight_layout()
    out = os.path.join(OUT_DIR, outname)
    fig.savefig(out, dpi=200)
    plt.close(fig)
    return out


def sig_mark(p):
    if not np.isfinite(p):
        return "n/a"
    if p < 0.001:
        return "***"
    if p < 0.01:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    per_image = collect()
    print(f"Collected {len(per_image)} images (5-slice subset)")
    slc = [d["n_slices"] for d in per_image]
    print(f"  slices/image used: min={min(slc)}, max={max(slc)} (target {N_SLICES})")
    for label, _, _ in CELL_LINES:
        c = {tp: len(vals(per_image, label, tp, "pct")) for tp in TIMEPOINTS}
        print(f"  {label}: " + ", ".join(f"{tp}min n={c[tp]}" for tp in TIMEPOINTS))

    metrics = [
        ("pct", "SG+ cells (%)", "Stress-granule-positive cells (%)  [5-slice subset]",
         "Fig_bar_SGpositive_percent.png", "Fig_line_SGpositive_percent.png"),
        ("mean_per_cell", "Number of SG per cell", "SG per cell (all cells)  [5-slice subset]",
         "Fig_bar_SG_per_cell.png", "Fig_line_SG_per_cell.png"),
        ("mean_per_pos", "SG per SG+ cell", "SG per SG+ cell  [5-slice subset]",
         "Fig_bar_SG_per_SGpositive.png", "Fig_line_SG_per_SGpositive.png"),
    ]
    figs = {}
    for key, ylabel, title, barname, linename in metrics:
        figs[(key, "bar")] = grouped_bar(per_image, key, ylabel, title + "  (mean +/- SD)", barname)
        figs[(key, "line")] = line_plot(per_image, key, ylabel, title + "  (mean +/- SEM)", linename)

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="404040")

    def style(ws, ncol, row=1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

    ws = wb.active
    ws.title = "Summary_per_image"
    # NB: the column that used to be headed "Slices_used" actually held the
    # slice COUNT; it is now "N_slices", and "Slices_used" lists the slices.
    cols = ["Cell_line", "Timepoint_min", "Image", "Source_CSV", "Source_image",
            "Acquisition_date", "Analysis_date", "Series", "Slices_used",
            "N_slices",
            "Total_cells", "Cells_with_punctae_(SG+)", "SG_positive_(%)",
            "Mean_punctae_per_cell", "Mean_punctae_per_SG+_cell", "Total_punctae"]
    ws.append(cols)
    order = {lab: i for i, (lab, _, _) in enumerate(CELL_LINES)}
    ordered = sorted(per_image,
                     key=lambda d: (order[d["line"]], d["tp"], d["series"]))
    for d in ordered:
        ws.append([d["line"], d["tp"], d["image"], d["csv"], d["source_image"],
                   d["acq_date"], d["ana_date"], d["series"], d["slices_str"],
                   d["n_slices"],
                   d["total_cells"], d["cells_with"], round(d["pct"], 3),
                   round(d["mean_per_cell"], 4), round(d["mean_per_pos"], 4),
                   d["total_punctae"]])
    style(ws, len(cols))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([12, 14, 46, 46, 46, 17, 15, 8, 16, 10,
                           12, 24, 14, 20, 22, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    write_sheet(
        wb.create_sheet("Per_cell_counts"),
        [("Cell_line", "line", 12), ("Timepoint_min", "tp", 14),
         ("Image", "image", 46), ("Source_CSV", "csv", 46),
         ("Source_image", "source_image", 46), ("Series", "series", 8)]
        + PER_CELL_TAIL,
        [dict(c, line=d["line"], tp=d["tp"], image=d["image"], csv=d["csv"],
              source_image=d["source_image"], series=d["series"],
              acq_date=d["acq_date"], ana_date=d["ana_date"])
         for d in ordered for c in d["cells"]])

    ws2 = wb.create_sheet("Group_means")
    ws2.append(["Cell_line", "Timepoint_min", "n_images",
                "Mean_SG+_(%)", "SD_SG+_(%)",
                "Mean_SG_per_cell", "SD_SG_per_cell",
                "Mean_SG_per_SG+cell", "SD_SG_per_SG+cell"])
    for label, _, _ in CELL_LINES:
        for tp in TIMEPOINTS:
            vp = vals(per_image, label, tp, "pct")
            if not vp:
                continue
            def ms(v):
                return (round(float(np.mean(v)), 4),
                        round(float(np.std(v, ddof=1)) if len(v) > 1 else 0.0, 4))
            m1, s1 = ms(vp)
            m2, s2 = ms(vals(per_image, label, tp, "mean_per_cell"))
            m3, s3 = ms(vals(per_image, label, tp, "mean_per_pos"))
            ws2.append([label, tp, len(vp), m1, s1, m2, s2, m3, s3])
    style(ws2, 9)
    for i, w in enumerate([12, 14, 10, 14, 14, 18, 18, 20, 20], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("Statistics")
    ws3.cell(1, 1, "One-way ANOVA across cell lines at each timepoint "
                   "(5-slice subset; per-image; n>=2 groups)").font = Font(bold=True)
    ws3.append(["Metric", "Timepoint_min", "cell_lines_(n)", "F", "p_value", "Significance"])
    style(ws3, 6, row=2)
    for key, ylabel, *_ in metrics:
        for tp in TIMEPOINTS:
            groups, tags = [], []
            for label, _, _ in CELL_LINES:
                v = vals(per_image, label, tp, key)
                if len(v) >= 2:
                    groups.append(np.array(v))
                    tags.append(f"{label}({len(v)})")
            if len(groups) >= 2:
                F, p = stats.f_oneway(*groups)
            else:
                F, p = float("nan"), float("nan")
            ws3.append([ylabel, tp, ", ".join(tags),
                        round(float(F), 4) if np.isfinite(F) else "n/a",
                        round(float(p), 6) if np.isfinite(p) else "n/a",
                        sig_mark(p)])
        ws3.append([])
    for i, w in enumerate([24, 14, 34, 12, 12, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("Figures")
    for (key, *_), row in zip(metrics, [1, 30, 59]):
        ws4.add_image(XLImage(figs[(key, "bar")]), f"A{row}")
        ws4.add_image(XLImage(figs[(key, "line")]), f"N{row}")

    out_xlsx = os.path.join(OUT_DIR, "ThreeCellLine_comparison_5slices.xlsx")
    wb.save(out_xlsx)

    print("\n=== Per-timepoint ANOVA across cell lines (5-slice subset) ===")
    for key, ylabel, *_ in metrics:
        line = []
        for tp in TIMEPOINTS:
            groups = [np.array(vals(per_image, lab, tp, key))
                      for lab, _, _ in CELL_LINES if len(vals(per_image, lab, tp, key)) >= 2]
            if len(groups) >= 2:
                F, p = stats.f_oneway(*groups)
                line.append(f"{tp}min p={p:.4g}({sig_mark(p)})")
        print(f"  {ylabel}: " + ", ".join(line))
    print("\nSaved to:", OUT_DIR)


if __name__ == "__main__":
    main()
