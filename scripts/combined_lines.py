"""
Combined multi-line SG+ time-course plot.

Plots SG-positive cells (%) vs timepoint with ONE line (mean +/- SD, with
individual image points) per user-defined group. Groups are matched by a
substring of the file name, so it works for cell lines, drugs, etc.

Usage:
    python combined_lines.py "PATH\\TO\\folder" "Label1:substr1" "Label2:substr2" ...

Example (parental vs R-KO As time course):
    python combined_lines.py "PATH" "RPMI:_RPMI_AS_" "RPMI R-KO:_RPMI_R_KO_AS_"

Each file is assigned to the FIRST label whose substring it contains, so list
the more specific substrings first if they overlap. Timepoint is read from
"<N>min" in the name; files without a timepoint (e.g. BTZ_3hr) are ignored.
Data rows are detected by an integer Cell_ID (works on .tif/.czi, skips SUMMARY).
Writes Figure_SGpositive_lines.png and SG_lines_summary.xlsx into the folder.
"""
import csv, glob, os, re, sys
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl

from sg_provenance import (PER_CELL_TAIL, cell_rows, parse_dates, slices_used,
                           write_sheet)

PALETTE = ["#4C72B0", "#C44E52", "#55A868", "#8172B2", "#CCB974", "#64B5CD"]
XLABEL = "Sodium arsenite (NaAsO2) exposure"


def is_data(r, h, ci):
    if len(r) < len(h):
        return False
    try:
        int(float(r[ci])); return True
    except (ValueError, IndexError):
        return False


def main(folder, specs):
    labels = [s.split(":", 1)[0] for s in specs]
    subs = [s.split(":", 1)[1] for s in specs]

    def group_of(fn):
        for lab, sub in zip(labels, subs):
            if sub in fn:
                return lab
        return None

    # (label, timepoint, SG+%, SG/cell, SG/SG+cell, image, cells, slices,
    #  n_slices, source_csv, acq_date, ana_date) -- vals() indexes 2/3/4.
    recs = []
    for path in sorted(glob.glob(os.path.join(folder, "*.csv"))):
        fn = os.path.basename(path)
        g = group_of(fn)
        tm = re.search(r"(\d+)\s*_?\s*min", fn)
        if g is None or tm is None:
            continue
        tp = int(tm.group(1))
        with open(path, newline="") as f:
            rows = list(csv.reader(f))
        h = rows[0]; ci = h.index("Cell_ID"); pi = h.index("Punctae_Count")
        si = h.index("Slice")
        data = [r for r in rows[1:] if is_data(r, h, ci)]
        if not data:
            continue
        # aggregate to unique Cell_IDs (sum punctae across slices); total cells =
        # number of unique Cell_IDs (Cell_ID 0 = background, excluded)
        cellp = {}
        cell_sl = {}
        for r in data:
            cid = r[ci].strip()
            if cid == "0":
                continue
            cellp[cid] = cellp.get(cid, 0) + int(float(r[pi]))
            cell_sl.setdefault(cid, set()).add(r[si].strip())
        if not cellp:
            continue
        ncell = len(cellp)
        withp = sum(1 for v in cellp.values() if v > 0)
        tot = sum(cellp.values())
        image = data[0][0]
        sl_str, n_sl = slices_used(cell_sl)
        ana_date, acq_date = parse_dates(fn)
        recs.append((g, tp, withp / ncell * 100, tot / ncell,
                     tot / withp if withp else 0.0, image,
                     cell_rows(cellp, cell_sl), sl_str, n_sl,
                     fn, acq_date, ana_date))

    timepoints = sorted({t for _, t, *_ in recs})

    def vals(lab, tp, idx):
        return np.array([r[idx] for r in recs if r[0] == lab and r[1] == tp])

    rng = np.random.default_rng(0)

    def plot_metric(idx, ylabel, title, outname):
        fig, ax = plt.subplots(figsize=(7, 5.5))
        x = np.arange(len(timepoints))
        for k, lab in enumerate(labels):
            color = PALETTE[k % len(PALETTE)]
            means = [vals(lab, tp, idx).mean() if len(vals(lab, tp, idx)) else np.nan for tp in timepoints]
            sds = [vals(lab, tp, idx).std(ddof=1) if len(vals(lab, tp, idx)) > 1 else 0.0 for tp in timepoints]
            ax.errorbar(x, means, yerr=sds, marker="o", capsize=4, lw=2, ms=7,
                        color=color, label=lab, zorder=3)
            for i, tp in enumerate(timepoints):
                v = vals(lab, tp, idx)
                ax.scatter(x[i] + rng.uniform(-0.08, 0.08, len(v)), v, s=14,
                           color=color, alpha=0.35, zorder=2)
        ax.set_xticks(x); ax.set_xticklabels([f"{t} min" for t in timepoints])
        ax.set_xlabel(XLABEL); ax.set_ylabel(ylabel); ax.set_title(title)
        ax.legend(title="Cell line", frameon=False); ax.set_ylim(bottom=0)
        fig.tight_layout()
        out = os.path.join(folder, outname)
        fig.savefig(out, dpi=200); plt.close(fig)
        return out

    out_png = plot_metric(2, "SG+ cells (%)", "Stress-granule-positive cells (%)",
                          "Figure_SGpositive_lines.png")
    plot_metric(3, "Number of SG per cell", "SG per cell",
                "Figure_SGperCell_lines.png")
    plot_metric(4, "Number of SG per SG+ cell", "SG per SG+ cell",
                "Figure_SGperSGposCell_lines.png")

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Per_image"
    ws.append(["Image", "Source_CSV", "Acquisition_date", "Analysis_date",
               "Group", "Timepoint_min", "Slices_used", "N_slices",
               "SG_positive_(%)", "SG_per_cell", "SG_per_SG+cell"])
    ordered = sorted(recs, key=lambda r: (labels.index(r[0]), r[1]))
    for r in ordered:
        ws.append([r[5], r[9], r[10], r[11], r[0], r[1], r[7], r[8],
                   round(r[2], 3), round(r[3], 4), round(r[4], 4)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    write_sheet(
        wb.create_sheet("Per_cell_counts"),
        [("Image", "image", 46), ("Source_CSV", "csv", 46),
         ("Group", "group", 16), ("Timepoint_min", "tp", 14)] + PER_CELL_TAIL,
        [dict(c, image=r[5], csv=r[9], group=r[0], tp=r[1],
              acq_date=r[10], ana_date=r[11])
         for r in ordered for c in r[6]])
    ws2 = wb.create_sheet("Group_means")
    ws2.append(["Group", "Timepoint_min", "n", "Mean_SG+_(%)", "SD_SG+_(%)",
                "Mean_SG_per_cell", "SD_SG_per_cell",
                "Mean_SG_per_SG+cell", "SD_SG_per_SG+cell"])
    for lab in labels:
        for tp in timepoints:
            n = len(vals(lab, tp, 2))
            def ms(idx):
                v = vals(lab, tp, idx)
                return (round(v.mean(), 4) if len(v) else None,
                        round(v.std(ddof=1), 4) if len(v) > 1 else 0.0)
            m2, s2 = ms(2); m3, s3 = ms(3); m4, s4 = ms(4)
            ws2.append([lab, tp, n, m2, s2, m3, s3, m4, s4])
    out_xlsx = os.path.join(folder, "SG_lines_summary.xlsx")
    wb.save(out_xlsx)

    for idx, name in ((2, "SG+ %"), (3, "SG/cell"), (4, "SG/SG+cell")):
        print(f"[{name}]")
        print(f"{'group':14}" + "".join(f"{str(t)+'min':>9}" for t in timepoints))
        for lab in labels:
            print(f"{lab:14}" + "".join(
                f"{(vals(lab,t,idx).mean() if len(vals(lab,t,idx)) else float('nan')):>9.2f}" for t in timepoints))
    print("Saved:", out_png)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit('Usage: python combined_lines.py "FOLDER" "Label1:substr1" "Label2:substr2" ...')
    main(sys.argv[1], sys.argv[2:])
